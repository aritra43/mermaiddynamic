from crewai import LLM, Agent, Crew, Process, Task
from crewai_tools import FileReadTool, FileWriterTool
from dotenv import load_dotenv
import os

load_dotenv()


llm = LLM(
    model="gemini/gemini-2.5-flash-preview-04-17"
    # model = "gemini/gemini-2.0-flash"
    # model="gpt-4o",
)

uploaded_srs="C:\\Users\\l43ar\\Downloads\\mermaiddynamic\\srs.docx"
def generate_sdd(uploaded_srs=uploaded_srs, original_filename="srs"):

        if uploaded_srs:
            print("[INFO] Upload received")
            filename = original_filename or "SRS.txt"
            name, ext = os.path.splitext(filename)
            ext = ext.lower()
            temp_dir = "../uploads"
            os.makedirs(temp_dir, exist_ok=True)

            # Normalized .txt file path
            srs_path = os.path.join(temp_dir, f"{name}.txt")

            try:
                if ext == ".docx":
                    import io
                    from docx import Document

                    print("[INFO] Converting DOCX to text...")
                    doc = Document(io.BytesIO(uploaded_srs))

                    text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
                    safe_text = ''.join([c if ord(c) < 128 or c in '\n\r\t' else '' for c in text])

                    with open(srs_path, "w", encoding="utf-8") as f:
                        f.write(safe_text)

                    print(f"[INFO] Cleaned DOCX saved to {srs_path}")

                elif ext == ".pdf":
                    import io
                    import fitz  # PyMuPDF

                    print("[INFO] Converting PDF to text...")
                    pdf = fitz.open(stream=io.BytesIO(uploaded_srs), filetype="pdf")
                    text = "\n".join([page.get_text() for page in pdf])
                    pdf.close()

                    safe_text = ''.join([c if ord(c) < 128 or c in '\n\r\t' else '' for c in text])

                    with open(srs_path, "w", encoding="utf-8") as f:
                        f.write(safe_text)

                    print(f"[INFO] Cleaned PDF saved to {srs_path}")

            except Exception as e:
                print(f"[ERROR] Failed to process uploaded file: {e}")
                raise

        file_read_tool = FileReadTool(file_path=srs_path)
        file_writer_tool = FileWriterTool()

        # Agents for SRS to SDD
        srs_extractor = Agent(
            role="SRS Extractor Agent",
            goal=(
                 "Extract structured information from the SRS document, categorizing key sections like Introduction, System Overview, "
                 "Functional and Non-Functional Requirements, API Design, Security & Compliance, and Data Management Strategy."
            ),
            backstory=(
                "A specialized AI agent trained in document analysis, natural language understanding, and structured data extraction. "
                "This agent identifies and categorizes key components from the SRS, ensuring no critical requirement is missed. "
                 "The extracted content serves as the foundation for subsequent tasks, including wireframe design and interface validation."
            ),
            llm=llm,
            tools=[file_read_tool]
        )

        sdd_structure = Agent(
            role="SDD Structure Agent",
            goal=(
                "Design a structured SDD template incorporating sections such as Introduction, System Overview, Non-Functional Requirements, "
                "API Design, Wireframe Designs, Interface Validation Rules, Security & Compliance, Data Management Strategy, and Appendices."
            ),
            backstory=(
                "A meticulous architect of technical documents, this agent ensures that the SDD follows best practices in software engineering. "
                "It constructs a logical and scalable framework that aligns with extracted requirements, ensuring completeness across all sections."
            ),
            llm=llm
        )

        er_schema_generator = Agent(
            role="ER Schema Generator Agent",
            goal=(
                "Extract the Data Model section from the SRS and generate a structured Entity Relationship (ER) schema with example database models."
            ),
            backstory=(
                 "A database design expert trained to analyze functional data models and convert them into structured ER schemas. "
                 "This agent ensures consistency, normalization, and correctness in entity definitions and relationships. "
                 "The ER Schema will be placed under the **System Architecture** section in the final SDD."
            ),
            llm=llm
        )
    

        content_generator = Agent(
            role="Content Generation Agent",
            goal=(
                "Populate the SDD template by mapping extracted SRS content to relevant sections, ensuring technical accuracy, coherence, and completeness."
            ),
            backstory=(
                 "An expert in technical writing and AI-driven content generation, this agent ensures clarity, precision, and completeness. "
                 "It generates high-quality descriptions, system overviews, non-functional requirements, API designs, security considerations, "
                 "wireframe descriptions, and interface validation rules. This agent ensures:"
                 "- **Introduction** includes Purpose, Scope, Assumptions, Constraints, and Stakeholders."
                 "- **System Overview** includes High-level Architecture (UML, DFD), Technology Stack, and Deployment Architecture."
                 "- **Non-Functional Requirements** include Performance, Scalability, Compliance, and Security Best Practices."
                 "- **ER Schema** includes example database models."
                 "- **API Design** includes Endpoints, Integration Standards, and Security & Authentication."
                 "- **Wireframe Designs** and **Interface Validation Rules** are included."
                 "- **Security & Compliance** includes Threat Modeling, Risk Assessment, Role-Based Access Control, and Data Management Strategy."
                 "- **Data Management Strategy** maps functional requirements to system components."
                 "- **Appendices** include Glossary and Compliance Standards."
            ),
            llm=llm
        )

        wireframe_designer = Agent(
            role="Wireframe Designer Agent",
            goal=(
                 "Create detailed descriptions of wireframes for each functional requirement in the SRS document."
            ),
            backstory=(
                "A skilled UI/UX designer trained in translating functional requirements into wireframe descriptions. "
                "This agent ensures that the wireframes align with the system's usability goals and user experience principles. "
                "The wireframe descriptions are structured and ready for integration into the SDD under the **Wireframe Designs** section."
            ),
            llm=llm
        )

        interface_validator = Agent(
            role="Interface Validation Agent",
            goal=(
                 "Define validation rules and user interface behavior for each functional requirement."
            ),
            backstory=(
                 "An expert in user interface design and validation, this agent ensures that all input fields, buttons, and interactions are properly validated and user-friendly. "
                 "It adheres to best practices in form validation, accessibility, and usability. "
                 "The validation rules and interaction guidelines are structured and ready for integration into the SDD under the **Interface Validation Rules** section."
            ),
            llm=llm
        )

        validation_compliance = Agent(
            role="Validation & Compliance Agent",
            goal=(
                 "Review and validate the SDD content to ensure alignment with security standards, regulatory compliance, and best practices."
            ),
            backstory=(
                "A compliance-focused AI agent trained in regulatory standards and best practices. "
                "It cross-verifies security implementations, role-based access controls, data management strategies, risk assessments, and API security. "
                "This agent ensures that all security & compliance documentation is properly included under Appendices."
            ),
            llm=llm
        )

        final_formatter = Agent(
            role="Final Formatting & Export Agent",
            goal=(
                "Format the finalized SDD content into professional output formats, ensuring readability, styling consistency, and professional presentation."
            ),
            backstory=(
                "A document refinement specialist with expertise in layout optimization and presentation. "
                "This agent transforms raw content into a polished, structured, and export-ready document (PDF, DOCX, Markdown), "
                "Ensuring all sections are complete and properly formatted."
            ),
            llm=llm,
            tools=[file_writer_tool]
        )

        # Define tasks
        extract_srs = Task(
            description=(
                "Analyze the SRS document and extract structured information, categorizing key sections like Introduction, System Overview, "
                "Functional and Non-Functional Requirements, API Design, Security & Compliance, and Data Management Strategy."
            ),
            expected_output=(
                 "A structured representation (JSON, dictionary, or tabular format) of the extracted SRS content, mapped to relevant SDD sections. "
                 "This extraction must explicitly include **Purpose, Scope, Assumptions, Constraints, Stakeholders**, **High-level Architecture**, "
                 "**Technology Stack**, **Deployment Architecture**, **Performance & Scalability**, **Compliance Standards**, "
                 "**Security Best Practices**, **ER Schema**, **API Endpoints**, **Wireframe Descriptions**, **Interface Validation Rules**, "
                 "**Threat Modeling**, **Risk Assessment**, **Role-Based Access Control**, **Data Management Strategy**, and **Glossary**."
            ),
            agent=srs_extractor
        )

        define_sdd_structure = Task(
            description=(
                "Design a structured SDD template incorporating sections such as Introduction, System Overview, Non-Functional Requirements, "
                "API Design, Wireframe Designs, Interface Validation Rules, Security & Compliance, Data Management Strategy, and Appendices."
            ),
            expected_output=(
                "A well-defined SDD template with placeholders for **Purpose, Scope, Assumptions, Constraints, Stakeholders**, "
                "**High-level Architecture (UML, DFD)**, **Technology Stack**, **Deployment Architecture**, **Performance & Scalability**, "
                "**Compliance Standards**, **Security Best Practices**, **ER Schema with Example Database Models**, **API Endpoints & Specifications**, "
                "**Integration Standards**, **Security & Authentication**, **Wireframe Designs**, **Interface Validation Rules**, "
                "**Threat Modeling**, **Risk Assessment**, **Role-Based Access Control**, **Data Management Strategy**, and **Appendices**."
            ),
            agent=sdd_structure
        )

        generate_er_schema = Task(
            description=(
                "Analyze the Data Model section in the Functional Requirements of the SRS document and generate a structured ER schema with example database models. "
                "The ER Schema must be placed under the **System Architecture** section in the final SDD."
            ),
            expected_output=(
                "A structured ER schema in text or JSON format representing all entities and relationships extracted from the Data Model section, with example database models."
            ),
            agent=er_schema_generator
        )

        generate_sdd_content = Task(
            description=(
                "Populate the SDD template by mapping extracted SRS content to relevant sections, ensuring technical accuracy, coherence, and completeness. "
                "Ensure the following are included: "
                "- **Introduction**: Purpose, Scope, Assumptions, Constraints, Stakeholders. "
                "- **System Overview**: High-level Architecture (UML, DFD), Technology Stack, Deployment Architecture. "
                "- **Non-Functional Requirements**: Performance, Scalability, Compliance Standards, Security Best Practices. "
                "- **ER Schema**: Example database models. "
                "- **API Design**: Endpoints, Integration Standards, Security & Authentication. "
                "- **Wireframe Designs**: Descriptions of wireframes for functional requirements. "
                "- **Interface Validation Rules**: Validation rules and interaction guidelines for functional requirements. "
                "- **Security & Compliance**: Threat Modeling, Risk Assessment, Role-Based Access Control, Data Management Strategy. "
                "- **Data Management Strategy**: Mapping of functional requirements to system components. "
                "- **Appendices**: Glossary, Compliance Standards."
                "-Wherever a diagram is needed, use the LLM-generated description to create code for the diagram using mermaidJS, following this format:\n"
                "mermaidstart\n"
                "    graph TD\n"
                "        A[Start] --> B{Decision}\n"
                "        B -- Yes --> C[Do Action 1]\n"
                "        B -- No --> D[Do Action 2]\n"
                "        C --> E[End]\n"
                "        D --> E[End]\n"
                "mermaidend"
                "For all the mermaid codes make it highly attractive and engaging."
                "In the code do not use () inside []"
                "Int styling do not use words such as User or Sequence after the word style"
            ),
            expected_output=(
                "A draft SDD document where all sections, including **Introduction**, **System Overview**, **Non-Functional Requirements**, **ER Schema**, "
                "**API Design**, **Wireframe Designs**, **Interface Validation Rules**, **Security & Compliance**, **Data Management Strategy**, and "
                "**Appendices**, are fully populated."
            ),
            agent=content_generator
        )

        generate_wireframe_descriptions = Task(
            description=(
                "Analyze the functional requirements from the SRS document and create detailed descriptions of wireframes aligned with each requirement."
            ),
            expected_output=(
                "Structured wireframe descriptions detailing the layout, elements, and interactions for each functional requirement."
            ),
            agent=wireframe_designer
        )

        define_interface_validation_rules = Task(
            description=(
               "Define validation rules and user interface behaviors for each functional requirement, ensuring usability and correctness."
            ),
            expected_output=(
                "Validation rules and interaction guidelines structured for integration into the SDD."
            ),
            agent=interface_validator
        )

        validate_sdd = Task(
            description=(
                "Review and validate the SDD content for security standards, regulatory compliance, and best practices."
            ),
            expected_output=(
                "A report detailing compliance status and suggestions for improvements in security, access control, data management, and regulatory adherence."
            ),
            agent=validation_compliance
        )

        format_export_sdd = Task(
            description=(
                "Format the finalized SDD content into professional output formats (PDF, DOCX, Markdown), ensuring proper layout, "
                "styling consistency, and readability. The formatted output should include all the populated sections from the SDD."
            ),
            expected_output=(
                f"A fully formatted and export-ready SDD document, ensuring all sections are complete and saved as `{original_filename}_sdd.md`."
            ),
            agent=final_formatter,
            output_file="{original_filename}_sdd.md"
        )

        # Initialize and execute the Crew
        crew = Crew(
            agents=[
                srs_extractor,
                sdd_structure,
                er_schema_generator,
                content_generator,
                wireframe_designer,
                interface_validator,
                validation_compliance,
                final_formatter
            ],
            tasks=[
                extract_srs,
                define_sdd_structure,
                generate_er_schema,
                generate_sdd_content,
                generate_wireframe_descriptions,
                define_interface_validation_rules,
                validate_sdd,
                format_export_sdd
            ],
            process=Process.sequential,
            verbose=True,
        )

        crew.kickoff()

        output_path = f"{original_filename}_sdd.md"

        if os.path.exists(output_path):
            with open(output_path, "r") as file:
                sdd_content = file.read()
            return sdd_content
        else:
            return "Error: File not found. Please Try Again!!"